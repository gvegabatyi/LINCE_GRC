"""
Módulo Exportar Informe — Generación de libro Excel consolidado.
Fase 2A: mantiene todas las hojas originales y añade de forma segura:
  - PERFIL_MPI
  - DELEGADO_DPD
  - RAT_DECRETO_662
  - REFERENCIAS_NORMATIVAS
"""

import streamlit as st
import pandas as pd
import io
from datetime import datetime
import json

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from db.repositories.derechos import get_solicitudes_df
from db.repositories.empresas import get_empresas_df
from db.repositories.riesgos import get_eipd_df, get_riesgos_df
from db.repositories.evaluaciones import get_evaluaciones_df
from db.repositories.tratamientos import get_tratamientos_df
from db.repositories.visitas import get_visitas_df
from db.repositories.evidencias import get_evidencias_df
from db.repositories.mpi_gobernanza import (
    get_perfiles_df,
    get_delegados_df,
    get_referencias_normativas_df,
)
from utils.constants import (
    CHECKLIST_PREGUNTAS,
    ESTADO_EVALUACION_DISPLAY,
    ESTADOS_MPI_INTERNOS_MAP,
    ESTADOS_MPI_OFICIALES_MAP,
    DERECHOS_TITULAR_MAP,
    ESTADOS_SOLICITUD_MAP,
    CATEGORIAS_RIESGO_MAP,
    GRAVEDAD_LEGAL_LABEL,
    RESPUESTAS_RIESGO_MAP,
    ESTADOS_RIESGO_MAP,
)
from services.riesgos_service import nivel_riesgo
from services import plazos_service as plazos
from services.checklist_service import calcular_porcentaje


# ==============================================================================
# ESTILOS EXCEL COMPARTIDOS
# ==============================================================================
def _estilos():
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font  = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    bold_font   = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    return header_fill, header_font, title_font, bold_font, regular_font, thin_border


def _format_sheet(ws, header_fill, header_font, thin_border, is_summary=False):
    """Aplica estilos a la primera fila de encabezados y ajusta anchos."""
    if not is_summary:
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 28

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            cell.border = thin_border
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)


# ==============================================================================
# GENERACIÓN DEL LIBRO EXCEL
# ==============================================================================
def generar_excel_completo() -> bytes:
    """
    Genera un archivo Excel consolidado en memoria con hojas formateadas.
    Retorna bytes del archivo .xlsx.
    """
    header_fill, header_font, title_font, bold_font, regular_font, thin_border = _estilos()

    wb = openpyxl.Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)

    # ── Hoja 1: EMPRESAS ──────────────────────────────────────────────────────
    ws1 = wb.create_sheet(title="EMPRESAS")
    df_emp = get_empresas_df()
    ws1.append(["ID", "Nombre Empresa", "RUT", "Dirección", "Contacto", "Email", "Teléfono", "Fecha Registro"])
    for _, r in df_emp.iterrows():
        ws1.append([
            r["id"], r["nombre"], r["rut"],
            r.get("direccion") or "", r.get("contacto") or "",
            r.get("email") or "", r.get("telefono") or "",
            str(r["created_at"])[:19] if pd.notnull(r.get("created_at")) else "",
        ])
    _format_sheet(ws1, header_fill, header_font, thin_border)

    # ── Hoja 2: TRATAMIENTOS ──────────────────────────────────────────────────
    ws2 = wb.create_sheet(title="TRATAMIENTOS")
    df_trat = get_tratamientos_df()
    ws2.append([
        "ID", "Empresa", "RUT Empresa", "Tratamiento", "Finalidades (v)",
        "Base Licitud (vii)", "Tipos de Datos (i)", "Titulares (iv)", "Evaluación", "Fecha Registro",
    ])
    for _, r in df_trat.iterrows():
        ws2.append([
            r["id"], r["empresa"], r["rut_empresa"], r["nombre_tratamiento"],
            r.get("rat_finalidades") or r.get("descripcion") or "",
            r.get("rat_base_licitud_interes_legitimo") or r.get("base_legal") or "",
            r.get("rat_tipo_datos_sensibles") or r.get("datos_personales") or "",
            r.get("rat_categorias_titulares") or "",
            r.get("evaluacion_nombre") or "",
            str(r["created_at"])[:19] if pd.notnull(r.get("created_at")) else "",
        ])
    _format_sheet(ws2, header_fill, header_font, thin_border)

    # ── Hoja 3: VISITAS ───────────────────────────────────────────────────────
    ws3 = wb.create_sheet(title="VISITAS")
    df_vis = get_visitas_df()
    headers_vis = [
        "ID", "Empresa", "RUT Empresa", "Fecha Visita", "Tipo Visita",
        "Auditor", "Puntaje (/15)", "% Cumplimiento", "Evaluación", "Observaciones",
    ]
    for q in CHECKLIST_PREGUNTAS:
        headers_vis.append(q[:30] + "…")
    ws3.append(headers_vis)

    for _, r in df_vis.iterrows():
        puntaje = int(r["puntaje"]) if pd.notnull(r.get("puntaje")) else 0
        porcentaje = f"{calcular_porcentaje(puntaje):.1f}%"
        row_data = [
            r["id"], r["empresa"], r["rut_empresa"], str(r["fecha_visita"]),
            r.get("tipo_visita") or "", r.get("auditor") or "",
            puntaje, porcentaje,
            r.get("evaluacion_nombre") or "", r.get("observaciones") or "",
        ]
        chk = r["checklist_seguridad"]
        if isinstance(chk, str):
            try:
                chk = json.loads(chk)
            except Exception:
                chk = {}
        elif not isinstance(chk, dict):
            chk = {}
        for q in CHECKLIST_PREGUNTAS:
            row_data.append(chk.get(q, "N/A"))
        ws3.append(row_data)
    _format_sheet(ws3, header_fill, header_font, thin_border)

    # ── Hoja 4: EVIDENCIAS ────────────────────────────────────────────────────
    ws4 = wb.create_sheet(title="EVIDENCIAS")
    df_ev = get_evidencias_df()
    ws4.append(["ID", "ID Visita", "Empresa", "Fecha Visita",
                "Descripción", "URL Foto / Archivo", "Evaluación", "Fecha Registro"])
    for _, r in df_ev.iterrows():
        ws4.append([
            r["id"], r["visita_id"], r["empresa"], str(r["fecha_visita"]),
            r.get("descripcion") or "", r.get("foto_url") or "",
            r.get("evaluacion_nombre") or "",
            str(r["created_at"])[:19] if pd.notnull(r.get("created_at")) else "",
        ])
    _format_sheet(ws4, header_fill, header_font, thin_border)

    # ── Hoja 5: RESUMEN ───────────────────────────────────────────────────────
    ws5 = wb.create_sheet(title="RESUMEN")
    ws5.append(["Lince GRC — Informe de Evaluación"])
    ws5.append(["Gobierno, Riesgo y Cumplimiento (Foco: Ley N° 21.719 / Decreto N° 662)"])
    ws5.append([f"Fecha de Generación: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"])
    ws5.append([])
    ws5.append(["Métrica General", "Valor"])
    ws5.append(["Total de Empresas Registradas", len(df_emp)])
    ws5.append(["Total de Tratamientos (RAT)", len(df_trat)])
    ws5.append(["Total de Visitas de Auditoría", len(df_vis)])
    ws5.append(["Total de Evidencias Adjuntas", len(df_ev)])

    promedio = df_vis["puntaje"].mean() if not df_vis.empty and "puntaje" in df_vis.columns else 0
    ws5.append(["Promedio Puntaje Seguridad (/15)", f"{promedio:.2f}"])
    ws5.append(["% Promedio Cumplimiento", f"{calcular_porcentaje(promedio):.1f}%"])

    ws5["A1"].font = title_font
    ws5["A2"].font = bold_font
    ws5["A3"].font = regular_font
    for cell in [ws5["A5"], ws5["B5"]]:
        cell.fill = header_fill
        cell.font = header_font
    for col in ws5.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws5.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 45)
        for cell in col:
            cell.border = thin_border

    # ── Hoja 6: EVALUACIONES ──────────────────────────────────────────────────
    try:
        df_eval = get_evaluaciones_df()
        ws6 = wb.create_sheet(title="EVALUACIONES")
        ws6.append([
            "ID", "Empresa", "RUT Empresa", "Nombre Evaluación", "Tipo",
            "Período", "Estado", "Auditor", "Fecha Inicio",
            "Cierre Planificado", "Cierre Real", "% Avance",
            "Alcance", "Objetivo", "Marco Normativo", "Observaciones", "Fecha Registro",
        ])
        for _, r in df_eval.iterrows():
            ws6.append([
                r["id"], r["empresa"], r["rut_empresa"], r["nombre"],
                r.get("tipo_evaluacion") or "",
                r.get("periodo") or "",
                ESTADO_EVALUACION_DISPLAY.get(r.get("estado", ""), r.get("estado", "")),
                r.get("auditor_responsable") or "",
                str(r["fecha_inicio"]) if pd.notnull(r.get("fecha_inicio")) else "",
                str(r["fecha_cierre_planificada"]) if pd.notnull(r.get("fecha_cierre_planificada")) else "",
                str(r["fecha_cierre_real"]) if pd.notnull(r.get("fecha_cierre_real")) else "",
                f"{float(r.get('porcentaje_avance') or 0):.1f}%",
                r.get("alcance") or "",
                r.get("objetivo") or "",
                r.get("marco_normativo") or "",
                r.get("observaciones") or "",
                str(r["created_at"])[:19] if pd.notnull(r.get("created_at")) else "",
            ])
        _format_sheet(ws6, header_fill, header_font, thin_border)
    except Exception:
        pass

    # ==========================================================================
    # HOJAS NUEVAS FASE 2A: DECRETO N° 662
    # ==========================================================================

    # ── Hoja 7: PERFIL_MPI ────────────────────────────────────────────────────
    try:
        df_perf = get_perfiles_df()
        ws7 = wb.create_sheet(title="PERFIL_MPI")
        ws7.append([
            "ID", "Empresa", "RUT Empresa", "Representante Legal (Art. 3 a)", "RUT Rep. Legal",
            "Fecha Adopción", "Versión", "Estado Interno", "Estado Registral Oficial",
            "Entidad Certificadora (Art. 15)", "N° Certificado", "Fecha Certificación",
            "Fecha Vencimiento", "N° Registro Agencia (Art. 17)", "Activo",
        ])
        for _, r in df_perf.iterrows():
            ws7.append([
                r["id"], r["empresa"], r["rut_empresa"],
                r["representante_legal_nombre"], r["representante_legal_rut"],
                str(r["fecha_adopcion_modelo"]) if pd.notnull(r.get("fecha_adopcion_modelo")) else "",
                r.get("version_modelo") or "1.0",
                ESTADOS_MPI_INTERNOS_MAP.get(r.get("estado_implementacion", ""), r.get("estado_implementacion", "")),
                ESTADOS_MPI_OFICIALES_MAP.get(r.get("estado_certificacion_oficial", ""), r.get("estado_certificacion_oficial", "")),
                r.get("entidad_certificadora") or "",
                r.get("nro_certificado") or "",
                str(r["fecha_certificacion"]) if pd.notnull(r.get("fecha_certificacion")) else "",
                str(r["fecha_vencimiento_certificacion"]) if pd.notnull(r.get("fecha_vencimiento_certificacion")) else "",
                r.get("nro_registro_agencia") or "",
                "Sí" if r.get("activo") else "No",
            ])
        _format_sheet(ws7, header_fill, header_font, thin_border)
    except Exception:
        pass

    # ── Hoja 8: DELEGADO_DPD ──────────────────────────────────────────────────
    try:
        df_del = get_delegados_df()
        ws8 = wb.create_sheet(title="DELEGADO_DPD")
        ws8.append([
            "ID", "Empresa", "Nombre Delegado (Art. 3 b)", "Correo Contacto Público (Art. 11)",
            "Teléfono", "Tipo Vínculo (Art. 7)", "Prestador Jurídico", "Cargo Interno",
            "Autoridad Designación (Art. 8)", "Fecha Designación", "Autonomía Declarada", "Activo",
        ])
        for _, r in df_del.iterrows():
            ws8.append([
                r["id"], r["empresa"], r["nombre_completo"],
                r["email_contacto_publico"], r.get("telefono_contacto") or "",
                r.get("tipo_vinculo") or "dependiente",
                r.get("persona_juridica_prestadora") or "",
                r.get("cargo_interno") or "",
                r.get("autoridad_designacion") or "Directorio",
                str(r["fecha_designacion"]) if pd.notnull(r.get("fecha_designacion")) else "",
                "Sí" if r.get("declaracion_autonomia") else "No",
                "Sí" if r.get("activo") else "No",
            ])
        _format_sheet(ws8, header_fill, header_font, thin_border)
    except Exception:
        pass

    # ── Hoja 9: RAT_DECRETO_662 (11 Numerales Art. 3 d) ───────────────────────
    try:
        ws9 = wb.create_sheet(title="RAT_DECRETO_662")
        ws9.append([
            "ID", "Empresa", "Tratamiento",
            "i. Tipos de Datos y Sensibles",
            "ii. Ámbito Territorial",
            "iii. Bases de Datos Administradas",
            "iv. Categorías de Titulares",
            "v. Finalidades del Tratamiento",
            "vi. Fuente de Datos / Acceso Público",
            "vii. Base Licitud / Interés Legítimo",
            "viii. Destinatarios / Cesiones",
            "ix. Transferencias Internacionales (Art. 27)",
            "x. Plazos / Supresión",
            "xi. Decisiones Automatizadas / Perfilamiento",
        ])
        for _, r in df_trat.iterrows():
            ws9.append([
                r["id"], r["empresa"], r["nombre_tratamiento"],
                r.get("rat_tipo_datos_sensibles") or "",
                r.get("rat_ambito_territorial") or "",
                r.get("rat_bases_datos_administradas") or "",
                r.get("rat_categorias_titulares") or "",
                r.get("rat_finalidades") or "",
                r.get("rat_fuente_datos_acceso_publico") or "",
                r.get("rat_base_licitud_interes_legitimo") or "",
                r.get("rat_destinatarios_cesiones") or "",
                r.get("rat_transferencias_internacionales_art27") or "",
                r.get("rat_plazo_condicion_supresion") or "",
                r.get("rat_decisiones_automatizadas_perfilamiento") or "",
            ])
        _format_sheet(ws9, header_fill, header_font, thin_border)
    except Exception:
        pass

    # ── Hoja 10: REFERENCIAS_NORMATIVAS ───────────────────────────────────────
    try:
        df_norm = get_referencias_normativas_df()
        ws10 = wb.create_sheet(title="REFERENCIAS_NORMATIVAS")
        ws10.append(["ID", "Fuente", "Artículo", "Materia", "Criterio Normativo"])
        if not df_norm.empty:
            for _, r in df_norm.iterrows():
                ws10.append([
                    r["id"], r.get("nombre_fuente") or "Decreto N° 662",
                    r["articulo"], r["materia"], r["descripcion_criterio"],
                ])
        _format_sheet(ws10, header_fill, header_font, thin_border)
    except Exception:
        pass

    # ==========================================================================
    # HOJAS NUEVAS FASE 2A.2: IMPORTACIONES DOCUMENTALES Y CANDIDATOS RAT
    # ==========================================================================

    # ── Hoja 11: IMPORTACIONES_DOCUMENTALES ───────────────────────────────────
    try:
        from db.repositories.importaciones import get_importaciones_df, get_candidatos_df
        import json as _json

        df_imp = get_importaciones_df()
        ws11 = wb.create_sheet(title="IMPORTACIONES_DOCUMENTALES")
        ws11.append([
            "ID", "Empresa", "Evaluación", "Nombre Archivo", "Tipo",
            "Tamaño (KB)", "Hash (abreviado)", "Estado",
            "Evidencia Permanente", "Total Candidatos",
            "Candidatos Creados", "Candidatos Ignorados",
            "Advertencias", "Fecha Importación",
        ])
        for _, r in df_imp.iterrows():
            ws11.append([
                r["id"],
                r.get("empresa") or "",
                r.get("evaluacion") or "",
                r.get("nombre_archivo") or "",
                r.get("tipo_archivo") or "",
                float(r.get("tamano_kb") or 0),
                r.get("hash_abreviado") or "",
                r.get("estado") or "",
                r.get("evidencia_permanente") or "No",
                int(r.get("total_candidatos") or 0),
                int(r.get("candidatos_creados") or 0),
                int(r.get("candidatos_ignorados") or 0),
                r.get("advertencias") or "",
                str(r["created_at"])[:19] if pd.notnull(r.get("created_at")) else "",
            ])
        _format_sheet(ws11, header_fill, header_font, thin_border)
    except Exception:
        pass

    # ── Hoja 12: CANDIDATOS_RAT ────────────────────────────────────────────────
    try:
        df_cand = get_candidatos_df()
        ws12 = wb.create_sheet(title="CANDIDATOS_RAT")
        ws12.append([
            "ID", "Empresa", "Archivo Origen", "Hoja/Tabla",
            "Fila", "Página", "Sección/Origen",
            "Estado Revisión", "Acción Humana",
            "Tratamiento Existente ID", "Tratamiento Creado ID",
            "Observación Revisión", "Fecha Revisión", "Fecha Registro",
        ])
        for _, r in df_cand.iterrows():
            ws12.append([
                r["id"],
                r.get("empresa") or "",
                r.get("nombre_archivo") or "",
                r.get("hoja") or "",
                r.get("nro_fila") or "",
                r.get("nro_pagina") or "",
                r.get("seccion_origen") or "",
                r.get("estado_revision") or "",
                r.get("accion_humana") or "",
                r.get("tratamiento_existente_id") or "",
                r.get("tratamiento_creado_id") or "",
                r.get("observacion_revision") or "",
                str(r["revisado_en"])[:19] if pd.notnull(r.get("revisado_en")) else "",
                str(r["created_at"])[:19] if pd.notnull(r.get("created_at")) else "",
            ])
        _format_sheet(ws12, header_fill, header_font, thin_border)
    except Exception:
        pass

    # ── Hoja 13: SOLICITUDES_DERECHOS (minimizada: sin nombre ni contacto del titular) ──
    try:
        df_sol = get_solicitudes_df()
        ws13 = wb.create_sheet(title="SOLICITUDES_DERECHOS")
        ws13.append([
            "Folio", "Empresa", "Derecho", "Régimen", "Fecha Ingreso", "Vencimiento",
            "Prorrogada", "Estado", "Estado del Plazo", "Bloqueo Solicitado",
            "Resultado Bloqueo", "Fecha Acuse", "Fecha Respuesta", "Identidad Verificada",
        ])
        for _, r in df_sol.iterrows():
            venc = pd.to_datetime(r["fecha_vencimiento"]).date()
            ws13.append([
                r["folio"],
                r["empresa"],
                DERECHOS_TITULAR_MAP.get(r["derecho"], r["derecho"]),
                "Ley 21.719" if r["regimen"] == plazos.REGIMEN_21719 else "Ley 19.628",
                str(r["fecha_ingreso"])[:10],
                str(venc),
                "Sí" if r["prorrogada"] else "No",
                ESTADOS_SOLICITUD_MAP.get(r["estado"], r["estado"]),
                plazos.estado_plazo(r["estado"], venc, r["regimen"]).etiqueta,
                "Sí" if r["bloqueo_solicitado"] else "No",
                r["bloqueo_resultado"] or "",
                str(r["fecha_acuse"])[:10] if pd.notnull(r["fecha_acuse"]) else "",
                str(r["fecha_respuesta"])[:10] if pd.notnull(r["fecha_respuesta"]) else "",
                "Sí" if r["identidad_verificada"] else "No",
            ])
        _format_sheet(ws13, header_fill, header_font, thin_border)
    except Exception:
        pass

    # ── Hoja 14: RIESGOS (registro y matriz) ─────────────────────────────────────
    try:
        df_rsg = get_riesgos_df()
        ws14 = wb.create_sheet(title="RIESGOS")
        ws14.append([
            "ID", "Empresa", "Tratamiento", "Riesgo", "Categoría", "Infracción (Ley 21.719)",
            "Gravedad Legal", "Probabilidad", "Impacto", "Nivel Inherente", "Controles",
            "Nivel Residual", "Respuesta", "Plan de Acción", "Responsable", "Plazo", "Estado", "Origen",
        ])
        for _, r in df_rsg.iterrows():
            _, _, niv = nivel_riesgo(r["probabilidad"], r["impacto"])
            residual = ""
            if pd.notnull(r["probabilidad_residual"]) and pd.notnull(r["impacto_residual"]):
                residual = nivel_riesgo(r["probabilidad_residual"], r["impacto_residual"])[2]
            ws14.append([
                r["id"], r["empresa"], r["tratamiento"] or "", r["titulo"],
                CATEGORIAS_RIESGO_MAP.get(r["categoria"], r["categoria"] or ""),
                r["infraccion_ref"] or "", GRAVEDAD_LEGAL_LABEL.get(r["gravedad_legal"], ""),
                int(r["probabilidad"]), int(r["impacto"]), niv, r["controles"] or "", residual,
                RESPUESTAS_RIESGO_MAP.get(r["respuesta"], "") if r["respuesta"] else "",
                r["plan_accion"] or "", r["responsable"] or "",
                str(r["plazo"])[:10] if pd.notnull(r["plazo"]) else "",
                ESTADOS_RIESGO_MAP.get(r["estado"], r["estado"]), r["origen"],
            ])
        _format_sheet(ws14, header_fill, header_font, thin_border)
    except Exception:
        pass

    # ── Hoja 15: EIPD (Ley 21.719 art. 15 ter) ─────────────────────────────────
    try:
        df_eipd = get_eipd_df()
        ws15 = wb.create_sheet(title="EIPD")
        ws15.append([
            "ID", "Empresa", "Tratamiento", "Fecha", "Disparador a", "Disparador b", "Disparador c",
            "Disparador d", "Conclusión", "Consulta a la Agencia", "Estado", "Responsable",
        ])
        for _, r in df_eipd.iterrows():
            ws15.append([
                r["id"], r["empresa"], r["tratamiento"] or "", str(r["fecha"])[:10],
                *("Sí" if r[f"disparador_{k}"] else "No" for k in "abcd"),
                {"riesgo_aceptable": "Riesgo residual aceptable", "alto_riesgo_residual": "Alto riesgo residual"}.get(r["conclusion"], ""),
                "Sí" if r["consulta_agencia"] else "No",
                "Aprobada" if r["estado"] == "aprobada" else "Borrador", r["responsable"] or "",
            ])
        _format_sheet(ws15, header_fill, header_font, thin_border)
    except Exception:
        pass

    _neutralizar_formulas(wb)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _neutralizar_formulas(wb) -> None:
    """
    Evita la inyección de fórmulas (CSV/Excel injection): el informe no usa
    fórmulas propias, así que todo texto que openpyxl interpretó como fórmula
    (empieza con '=') proviene de datos del usuario y se fuerza a texto literal.
    """
    for ws in wb.worksheets:
        for fila in ws.iter_rows():
            for celda in fila:
                if celda.data_type == "f":
                    celda.data_type = "s"


# ==============================================================================
# MÓDULO UI
# ==============================================================================
def render() -> None:
    st.header("Generación y Descarga de Informe Completo")
    st.write("Genera un libro Excel consolidado (`.xlsx`) con hojas formateadas profesionalmente:")

    st.markdown("""
    **Hojas de Gestión General:**
    - **Hoja 1 — EMPRESAS:** Catálogo de empresas del grupo.
    - **Hoja 2 — TRATAMIENTOS:** Resumen ejecutivo de tratamientos.
    - **Hoja 3 — VISITAS:** Auditorías con puntaje, tipo y desglose del checklist de seguridad.
    - **Hoja 4 — EVIDENCIAS:** Registro y enlaces de evidencias levantadas en terreno.
    - **Hoja 5 — RESUMEN:** Métricas ejecutivas y fecha de emisión.
    - **Hoja 6 — EVALUACIONES:** Proyectos GRC por empresa y período.

    **Hojas del Modelo de Prevención de Infracciones (art. 49 Ley 21.719 · Decreto N° 662):**
    - **Hoja 7 — PERFIL_MPI:** Responsable, representante legal, estados interno y oficial *(Art. 3 a; Arts. 15 y 17)*.
    - **Hoja 8 — DELEGADO_DPD:** Delegado DPD, canal de contacto accesible e idoneidad *(Arts. 3 b, c; Arts. 6-13)*.
    - **Hoja 9 — RAT_DECRETO_662:** Registro de Actividades con los **11 numerales taxativos** *(Art. 3 d)*.
    - **Hoja 10 — REFERENCIAS_NORMATIVAS:** Citas y criterios del Decreto N° 662.

    **Hojas de Importación Documental:**
    - **Hoja 11 — IMPORTACIONES_DOCUMENTALES:** Historial de archivos importados, hash, estado y candidatos generados.
    - **Hoja 12 — CANDIDATOS_RAT:** Candidatos extraídos, estado de revisión y acción humana aplicada.

    **Derechos de los Titulares:**
    - **Hoja 13 — SOLICITUDES_DERECHOS:** Folio, derecho, régimen, plazos y estado. Por minimización no incluye nombre ni contacto del titular.

    **Análisis de Riesgos:**
    - **Hoja 14 — RIESGOS:** Registro de riesgos con infracción asociada, niveles inherente y residual y plan de acción.
    - **Hoja 15 — EIPD:** Evaluaciones de impacto por tratamiento (art. 15 ter): disparadores, conclusión y estado.
    """)

    st.markdown("---")

    col_exp1, col_exp2 = st.columns([1, 2])
    with col_exp1:
        if st.button("Generar Informe Excel", type="primary", use_container_width=True):
            with st.spinner("Generando libro Excel…"):
                try:
                    excel_bytes = generar_excel_completo()
                    st.session_state["excel_informe"] = excel_bytes
                    st.success("¡Informe consolidado generado con éxito!")
                except Exception as exc:
                    st.error(f"Error al generar el informe ({type(exc).__name__}). Intente nuevamente.")

    with col_exp2:
        if "excel_informe" in st.session_state:
            nombre_archivo = f"Lince_GRC_Informe_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            st.download_button(
                label="Descargar Informe Completo Excel (.xlsx)",
                data=st.session_state["excel_informe"],
                file_name=nombre_archivo,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )
