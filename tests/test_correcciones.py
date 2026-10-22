"""
Pruebas de regresión de las correcciones de la auditoría (docs/AUDITORIA_ESTADO_PROYECTO.md §12).

- Unitarias: sin base de datos.
- Integración: sólo si se define LINCE_TEST_DATABASE_URL apuntando a una BD
  DESECHABLE (nunca a la de desarrollo o producción). Ejemplo:
      docker run -d --name lince-test -e POSTGRES_PASSWORD=test -p 127.0.0.1:55432:5432 postgres:16-alpine
      LINCE_TEST_DATABASE_URL=postgresql://postgres:test@localhost:55432/postgres \
          python -m unittest discover -s tests -t .
"""

import csv
import io
import os
import unittest
from unittest import mock

import openpyxl

from db.repositories.importaciones import _sanitizar_matriz
from modules.exportar import _neutralizar_formulas
from services.extraccion_service import generar_candidatos_csv, generar_candidatos_xlsx


def _xlsx(filas: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RAT"
    for f in filas:
        ws.append(f)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _csv(filas: list[list]) -> bytes:
    buf = io.StringIO()
    csv.writer(buf).writerows(filas)
    return buf.getvalue().encode("utf-8")


FILAS = [
    ["Tratamiento", "Finalidad"],
    ["Nómina", "Pago de remuneraciones"],
    ["Clientes", "Facturación"],
    ["Proveedores", "Compras"],
]
MAPEO = {0: "nombre_tratamiento", 1: "rat_finalidades"}


class TestImportacionXlsx(unittest.TestCase):
    def test_conserva_primera_fila_de_datos(self):
        candidatos, _, utiles = generar_candidatos_xlsx(_xlsx(FILAS), "RAT", MAPEO)
        self.assertEqual(utiles, 3)
        self.assertEqual(_nombres(candidatos), ["Nómina", "Clientes", "Proveedores"])
        # nro_fila = fila real en Excel (fila 1 = encabezado)
        self.assertEqual([c["nro_fila"] for c in candidatos], [2, 3, 4])

    def test_paridad_con_csv(self):
        cand_x, _, _ = generar_candidatos_xlsx(_xlsx(FILAS), "RAT", MAPEO)
        cand_c, _, _ = generar_candidatos_csv(_csv(FILAS), MAPEO)
        self.assertEqual(_nombres(cand_x), _nombres(cand_c))
        self.assertEqual([c["nro_fila"] for c in cand_x], [c["nro_fila"] for c in cand_c])


def _nombres(candidatos: list[dict]) -> list[str]:
    return [c["datos_extraidos_json"]["nombre_tratamiento"] for c in candidatos]


class TestExportacion(unittest.TestCase):
    def test_neutraliza_formulas(self):
        wb = openpyxl.Workbook()
        wb.active.append(['=HYPERLINK("http://x","clic")', "texto", 42])
        _neutralizar_formulas(wb)
        buf = io.BytesIO()
        wb.save(buf)
        ws = openpyxl.load_workbook(io.BytesIO(buf.getvalue())).active
        self.assertEqual(ws["A1"].data_type, "s")
        self.assertEqual(ws["A1"].value, '=HYPERLINK("http://x","clic")')
        self.assertEqual(ws["C1"].value, 42)


class TestPersistenciaImportacion(unittest.TestCase):
    def test_trunca_celdas_a_4000(self):
        (fila,) = _sanitizar_matriz([["x" * 5000, None, 3]])
        self.assertEqual(len(fila[0]), 4000)
        self.assertEqual(fila[1:], ["", "3"])


class TestMigracionPaso8(unittest.TestCase):
    def test_identificadores_fijos_y_seguros(self):
        from db.migrations import _FK_CORRECTIVAS
        for ident in (x for fk in _FK_CORRECTIVAS for x in fk):
            self.assertRegex(ident, r"^[a-z_]+$")


@unittest.skipUnless(os.getenv("LINCE_TEST_DATABASE_URL"), "Definir LINCE_TEST_DATABASE_URL (BD desechable)")
class TestIntegracionBD(unittest.TestCase):
    """Ejecuta init_db() real sobre una BD desechable y verifica el esquema resultante."""

    @classmethod
    def setUpClass(cls):
        cls.url = os.environ["LINCE_TEST_DATABASE_URL"]
        cls.env = mock.patch.dict(os.environ, {"DATABASE_URL": cls.url})
        cls.env.start()
        from db.migrations import init_db
        init_db()
        init_db()  # idempotencia: la segunda ejecución no debe fallar ni duplicar

    @classmethod
    def tearDownClass(cls):
        cls.env.stop()

    def _query(self, sql, params=None):
        from db.connection import get_connection
        conn = get_connection()
        try:
            with conn.cursor() as c:
                c.execute(sql, params)
                return c.fetchall()
        finally:
            conn.close()

    def test_tablas(self):
        tablas = {r[0] for r in self._query(
            "SELECT table_name FROM information_schema.tables WHERE table_schema='public'")}
        esperadas = {"empresas", "tratamientos", "visitas", "evidencias", "evaluaciones",
                     "normativa_fuentes", "normativa_referencias", "mpi_perfiles", "mpi_delegados",
                     "importaciones_documentales", "importacion_tratamientos_candidatos"}
        self.assertTrue(esperadas <= tablas, esperadas - tablas)

    def test_fk_correctivas(self):
        from db.migrations import _FK_CORRECTIVAS
        for tabla, columna, referida in _FK_CORRECTIVAS:
            filas = self._query(
                """
                SELECT ccu.table_name, rc.delete_rule
                FROM information_schema.table_constraints tc
                JOIN information_schema.key_column_usage kcu
                  ON tc.constraint_name = kcu.constraint_name AND tc.table_schema = kcu.table_schema
                JOIN information_schema.referential_constraints rc
                  ON rc.constraint_name = tc.constraint_name AND rc.constraint_schema = tc.table_schema
                JOIN information_schema.constraint_column_usage ccu
                  ON ccu.constraint_name = tc.constraint_name AND ccu.table_schema = tc.table_schema
                WHERE tc.constraint_type='FOREIGN KEY' AND tc.table_name=%s AND kcu.column_name=%s
                """,
                (tabla, columna),
            )
            self.assertEqual(len(filas), 1, f"{tabla}.{columna}: {filas}")
            self.assertEqual(filas[0], (referida, "SET NULL"))

    def test_ambito_territorial_es_text(self):
        (tipo,), = self._query(
            "SELECT data_type FROM information_schema.columns "
            "WHERE table_name='tratamientos' AND column_name='rat_ambito_territorial'")
        self.assertEqual(tipo, "text")

    def test_referencia_colgante_se_anula(self):
        """Simula una BD antigua sin la FK y con un ID colgante: el Paso 8 debe anularlo y crear la FK."""
        from db.connection import get_connection
        from db.migrations import init_db
        conn = get_connection()
        try:
            with conn.cursor() as c:
                c.execute("ALTER TABLE tratamientos DROP CONSTRAINT IF EXISTS fk_tratamientos_evaluacion_id")
                c.execute("INSERT INTO empresas (nombre, rut) VALUES ('Test colgante', 'TEST-COLG') RETURNING id")
                empresa_id = c.fetchone()[0]
                c.execute(
                    "INSERT INTO tratamientos (empresa_id, nombre_tratamiento, evaluacion_id) "
                    "VALUES (%s, 'T colgante', 999999) RETURNING id", (empresa_id,))
                trat_id = c.fetchone()[0]
            conn.commit()
        finally:
            conn.close()
        try:
            init_db()
            (evaluacion,), = self._query("SELECT evaluacion_id FROM tratamientos WHERE id=%s", (trat_id,))
            self.assertIsNone(evaluacion)
            self.test_fk_correctivas()
        finally:
            conn = get_connection()
            with conn.cursor() as c:
                c.execute("DELETE FROM empresas WHERE id=%s", (empresa_id,))
            conn.commit()
            conn.close()

    def test_seed_normativo(self):
        (n,), = self._query("SELECT count(*) FROM normativa_referencias")
        self.assertEqual(n, 21)


if __name__ == "__main__":
    unittest.main()
