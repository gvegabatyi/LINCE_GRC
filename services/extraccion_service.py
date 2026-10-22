"""
Servicio de Extracción Determinista — Fase 2A.2 (Lince GRC).

Extrae texto y candidatos de campos RAT a partir de documentos en memoria.
Soporta: XLSX, CSV, DOCX, PDF (con texto), MD, TXT.

Principios:
- Sin IA externa, sin OCR, sin inferencia automática de campos.
- No almacena el archivo original. La importación guarda en la base de datos sólo las
  columnas mapeadas (las demás vacías, salvo opt-in), con celdas truncadas a 4.000 caracteres.
- Toda información extraída es un borrador no validado.
- Los datos definitivos solo se persisten por acción humana explícita.
- Límites de privacidad: máx. 4.000 caracteres por campo (truncamiento explícito con advertencia).
- Campos no mapeados NO se almacenan salvo opt-in explícito del usuario.
- Resumen de extracción almacena exclusivamente metadatos técnicos (sin datos personales ni valores de filas).
- Procesamiento de filas útiles: no se descarta ninguna fila útil con datos por campos incompletos.
- Regla de combinación: columnas con el mismo destino se combinan con ' / ' sin sobrescritura destructiva.
- Nombre de tratamiento por defecto: 'Nombre de tratamiento pendiente de identificación'.
- Previsualización completa: expone todas las filas de datos para navegación interactiva sin recorte a 5 filas.
"""

import io
import json


# ==============================================================================
# CAMPOS RAT DISPONIBLES PARA MAPEO MANUAL
# Art. 3 letra d) Decreto N° 662, numerales i a xi + campos GRC complementarios
# ==============================================================================

CAMPOS_RAT_DECRETO = [
    ("nombre_tratamiento",                          "Nombre del tratamiento"),
    ("descripcion",                                 "Descripción general"),
    ("rat_tipo_datos_sensibles",                    "i. Tipos de datos y datos sensibles (Art. 3 d-i)"),
    ("rat_ambito_territorial",                      "ii. Ámbito territorial (Art. 3 d-ii)"),
    ("rat_bases_datos_administradas",               "iii. Bases de datos administradas (Art. 3 d-iii)"),
    ("rat_categorias_titulares",                    "iv. Categorías de titulares (Art. 3 d-iv)"),
    ("rat_finalidades",                             "v. Finalidades del tratamiento (Art. 3 d-v)"),
    ("rat_fuente_datos_acceso_publico",             "vi. Fuente de datos / acceso público (Art. 3 d-vi)"),
    ("rat_base_licitud_interes_legitimo",           "vii. Base de licitud / interés legítimo (Art. 3 d-vii)"),
    ("rat_destinatarios_cesiones",                  "viii. Destinatarios y cesiones (Art. 3 d-viii)"),
    ("rat_transferencias_internacionales_art27",    "ix. Transferencias internacionales Art. 27 (Art. 3 d-ix)"),
    ("rat_plazo_condicion_supresion",               "x. Plazo / condición de supresión (Art. 3 d-x)"),
    ("rat_decisiones_automatizadas_perfilamiento",  "xi. Decisiones automatizadas / perfilamiento (Art. 3 d-xi)"),
    # Campos complementarios GRC
    ("base_legal",      "Base legal (complementario GRC)"),
    ("datos_personales","Datos personales (complementario GRC)"),
    ("responsables",    "Responsables (complementario GRC)"),
    ("encargados",      "Encargados del tratamiento (complementario GRC)"),
    ("transferencias",  "Transferencias (complementario GRC)"),
]

CAMPOS_RAT_CLAVES = [c[0] for c in CAMPOS_RAT_DECRETO]
CAMPOS_RAT_ETIQUETAS = {c[0]: c[1] for c in CAMPOS_RAT_DECRETO}

MAX_CANDIDATOS_POR_VISTA = 100
MAX_CARACTERES_CAMPO = 4000
NOMBRE_TRATAMIENTO_DEFECTO = "Nombre de tratamiento pendiente de identificación"


# ==============================================================================
# UTILIDAD DE TRUNCAMIENTO SEGURO
# ==============================================================================

def limitar_texto_campo(val: str, max_chars: int = MAX_CARACTERES_CAMPO) -> tuple[str, bool]:
    """
    Limita el texto a max_chars caracteres (4.000 por defecto).
    Retorna (texto_limpio_o_truncado, fue_truncado).
    """
    if not val:
        return "", False
    val_str = str(val).strip()
    if len(val_str) > max_chars:
        return val_str[:max_chars], True
    return val_str, False


def _combinar_valores_campo(existente: str, nuevo: str) -> str:
    """
    Combina dos valores mapeados al mismo campo destino de forma determinista y segura:
    - Si el existente está vacío, toma el nuevo.
    - Si el nuevo está vacío, conserva el existente.
    - Si ambos tienen texto y son idénticos, no duplica.
    - Si ambos tienen texto distinto, combina con ' / '.
    """
    ex = existente.strip() if existente else ""
    nv = nuevo.strip() if nuevo else ""
    if not ex:
        return nv
    if not nv:
        return ex
    if ex == nv or nv in ex.split(" / "):
        return ex
    return f"{ex} / {nv}"


def _es_fila_util(celdas: list[str], encabezados: list[str] | None = None) -> bool:
    """
    Determina si una fila es útil:
    - Falsa si todas las celdas están vacías.
    - Falsa si todas las celdas coinciden literalmente con la fila de encabezados.
    - Verdadera si contiene al menos una celda con texto no vacío.
    """
    celdas_limpias = [str(c).strip() for c in celdas if c is not None]
    if not any(c for c in celdas_limpias):
        return False

    if encabezados:
        enc_limpios = [str(h).strip().lower() for h in encabezados if h is not None]
        fila_comparar = [c.lower() for c in celdas_limpias]
        if fila_comparar == enc_limpios:
            return False

    return True


# ==============================================================================
# XLSX
# ==============================================================================

def extraer_xlsx_hojas(contenido: bytes) -> list[str]:
    """Retorna los nombres de las hojas disponibles en el XLSX."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        return wb.sheetnames
    except Exception as exc:
        raise ValueError(f"No se pudo leer el archivo XLSX: {exc}")


def extraer_xlsx_encabezados_preview(contenido: bytes, nombre_hoja: str) -> tuple[list[str], list[list], int]:
    """
    Retorna (encabezados, todas_filas_datos, total_filas_datos) de la hoja seleccionada.
    """
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        ws = wb[nombre_hoja]
        filas = []
        for row in ws.iter_rows(values_only=True):
            filas.append([str(v) if v is not None else "" for v in row])
        if not filas:
            return [], [], 0
        encabezados = filas[0]
        datos = filas[1:]
        return encabezados, datos, len(datos)
    except Exception as exc:
        raise ValueError(f"No se pudo leer la hoja '{nombre_hoja}': {exc}")


def generar_candidatos_xlsx(
    contenido: bytes,
    nombre_hoja: str,
    mapeo_columnas: dict,
    fila_inicio: int = 0,  # índice sobre las filas de datos (el encabezado ya está excluido)
    limite: int = MAX_CANDIDATOS_POR_VISTA,
    conservar_adicionales: bool = False,
) -> tuple[list[dict], int, int]:
    """
    Genera candidatos desde filas XLSX.
    Retorna (candidatos, total_truncados, total_filas_utiles).
    """
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        ws = wb[nombre_hoja]
        filas = list(ws.iter_rows(values_only=True))
    except Exception as exc:
        raise ValueError(f"Error al leer filas del XLSX: {exc}")

    if not filas:
        return [], 0, 0

    encabezados = [str(v).strip() if v is not None else "" for v in filas[0]]
    datos_filas = filas[1:] if len(filas) > 1 else []

    candidatos = []
    total_truncados = 0
    total_filas_utiles = 0

    for idx_relativo, fila_raw in enumerate(datos_filas[fila_inicio:], start=fila_inicio):
        fila = [str(v).strip() if v is not None else "" for v in fila_raw]
        if not _es_fila_util(fila, encabezados):
            continue

        total_filas_utiles += 1
        if len(candidatos) >= limite:
            continue

        nro_fila_original = idx_relativo + 2
        campos_candidato = {}
        campos_adicionales = {}

        for idx_col, valor in enumerate(fila):
            val_limpio, fue_truncado = limitar_texto_campo(valor)
            if fue_truncado:
                total_truncados += 1

            campo_destino = mapeo_columnas.get(idx_col)
            if campo_destino:
                campos_candidato[campo_destino] = _combinar_valores_campo(
                    campos_candidato.get(campo_destino, ""), val_limpio
                )
            elif conservar_adicionales and val_limpio:
                campos_adicionales[f"col_{idx_col}"] = val_limpio

        # Nombre por defecto si la fila útil no contiene nombre
        if not campos_candidato.get("nombre_tratamiento", "").strip():
            campos_candidato["nombre_tratamiento"] = NOMBRE_TRATAMIENTO_DEFECTO

        datos_json = dict(campos_candidato)
        if conservar_adicionales and campos_adicionales:
            datos_json["campos_adicionales"] = campos_adicionales

        candidatos.append({
            "nro_hoja": nombre_hoja,
            "nro_fila": nro_fila_original,
            "nro_pagina": None,
            "seccion_origen": None,
            "datos_extraidos_json": datos_json,
        })

    return candidatos, total_truncados, total_filas_utiles


# ==============================================================================
# CSV
# ==============================================================================

def extraer_csv_encabezados_preview(contenido: bytes) -> tuple[list[str], list[list], str, int]:
    """
    Detecta delimitador y retorna (encabezados, todas_filas_datos, delimitador, total_filas_datos).
    """
    import csv

    texto = _decodificar_texto(contenido)
    sample = texto[:4096]

    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimitador = dialect.delimiter
    except Exception:
        delimitador = ","

    reader = csv.reader(io.StringIO(texto), delimiter=delimitador)
    filas = list(reader)
    if not filas:
        return [], [], delimitador, 0

    encabezados = [c.strip() for c in filas[0]]
    datos = [r for r in filas[1:]]
    return encabezados, datos, delimitador, len(datos)


def generar_candidatos_csv(
    contenido: bytes,
    mapeo_columnas: dict,
    fila_inicio: int = 0,
    limite: int = MAX_CANDIDATOS_POR_VISTA,
    conservar_adicionales: bool = False,
) -> tuple[list[dict], int, int]:
    """Genera candidatos desde filas CSV."""
    import csv
    texto = _decodificar_texto(contenido)

    sample = texto[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimitador = dialect.delimiter
    except Exception:
        delimitador = ","

    reader = csv.reader(io.StringIO(texto), delimiter=delimitador)
    filas = list(reader)
    if not filas:
        return [], 0, 0

    encabezados = [c.strip() for c in filas[0]]
    datos_filas = filas[1:] if len(filas) > 1 else []

    candidatos = []
    total_truncados = 0
    total_filas_utiles = 0

    for idx_relativo, fila in enumerate(datos_filas[fila_inicio:], start=fila_inicio):
        fila_limpia = [c.strip() for c in fila]
        if not _es_fila_util(fila_limpia, encabezados):
            continue

        total_filas_utiles += 1
        if len(candidatos) >= limite:
            continue

        nro_fila_original = idx_relativo + 2
        campos_candidato = {}
        campos_adicionales = {}

        for idx_col, valor in enumerate(fila_limpia):
            val_limpio, fue_truncado = limitar_texto_campo(valor)
            if fue_truncado:
                total_truncados += 1

            campo_destino = mapeo_columnas.get(idx_col)
            if campo_destino:
                campos_candidato[campo_destino] = _combinar_valores_campo(
                    campos_candidato.get(campo_destino, ""), val_limpio
                )
            elif conservar_adicionales and val_limpio:
                campos_adicionales[f"col_{idx_col}"] = val_limpio

        if not campos_candidato.get("nombre_tratamiento", "").strip():
            campos_candidato["nombre_tratamiento"] = NOMBRE_TRATAMIENTO_DEFECTO

        datos_json = dict(campos_candidato)
        if conservar_adicionales and campos_adicionales:
            datos_json["campos_adicionales"] = campos_adicionales

        candidatos.append({
            "nro_hoja": "CSV",
            "nro_fila": nro_fila_original,
            "nro_pagina": None,
            "seccion_origen": None,
            "datos_extraidos_json": datos_json,
        })

    return candidatos, total_truncados, total_filas_utiles


# ==============================================================================
# DOCX (TABLAS Y SECCIONES DE TEXTO)
# ==============================================================================

def extraer_docx_estructura(contenido: bytes) -> dict:
    """
    Extrae la estructura completa del DOCX: tablas y secciones de párrafos.
    Para cada tabla reporta:
    - encabezados reales (fila 0)
    - filas_datos completas (todas las filas de datos)
    - filas_preview (primeras 5 filas de muestra)
    - total_filas (total de filas de datos detectadas)
    - filas_texto completas
    """
    try:
        from docx import Document
    except ImportError:
        raise ImportError("La biblioteca 'python-docx' no está instalada. Ejecute: pip install python-docx")

    try:
        doc = Document(io.BytesIO(contenido))
    except Exception as exc:
        raise ValueError(f"No se pudo abrir el archivo DOCX: {exc}")

    tablas = []
    for idx_tabla, tabla in enumerate(doc.tables):
        try:
            filas_texto = []
            for fila in tabla.rows:
                celdas = [cell.text.strip() for cell in fila.cells]
                filas_texto.append(celdas)

            encabezados = filas_texto[0] if filas_texto else []
            filas_datos = filas_texto[1:] if len(filas_texto) > 1 else []
            preview = filas_datos[:5]
            total_datos = len(filas_datos)

            tablas.append({
                "indice": idx_tabla,
                "encabezados": encabezados,
                "filas_datos": filas_datos,
                "filas_preview": preview,
                "total_filas": total_datos,
                "filas_texto": filas_texto,
            })
        except Exception:
            tablas.append({
                "indice": idx_tabla,
                "encabezados": [],
                "filas_datos": [],
                "filas_preview": [],
                "total_filas": 0,
                "filas_texto": [],
                "advertencia": "No se pudo leer esta tabla.",
            })

    secciones = []
    seccion_actual = {"titulo": "(Sin título)", "parrafos": []}
    for para in doc.paragraphs:
        if para.style.name.startswith("Heading"):
            if seccion_actual["parrafos"]:
                secciones.append({
                    "titulo": seccion_actual["titulo"],
                    "texto_preview": "\n".join(seccion_actual["parrafos"])[:500],
                    "texto_completo": "\n".join(seccion_actual["parrafos"]),
                })
            seccion_actual = {"titulo": para.text.strip() or "(Sin título)", "parrafos": []}
        else:
            if para.text.strip():
                seccion_actual["parrafos"].append(para.text.strip())

    if seccion_actual["parrafos"]:
        secciones.append({
            "titulo": seccion_actual["titulo"],
            "texto_preview": "\n".join(seccion_actual["parrafos"])[:500],
            "texto_completo": "\n".join(seccion_actual["parrafos"]),
        })

    return {"tablas": tablas, "secciones": secciones}


def generar_candidatos_docx_tabla(
    estructura: dict,
    indice_tabla: int,
    mapeo_columnas: dict,
    limite: int = MAX_CANDIDATOS_POR_VISTA,
    conservar_adicionales: bool = False,
) -> tuple[list[dict], int, int, int]:
    """
    Genera candidatos desde una tabla específica del DOCX.

    REGLAS DE ROBUSTEZ:
    1. Procesa cada fila útil (excluye únicamente filas vacías o repeticiones literales de encabezados).
    2. Si una tabla tiene 12 filas reales de datos, genera exactamente 12 candidatos.
    3. Si varias columnas se mapean al mismo campo destino, combina con ' / ' sin borrar valores no vacíos.
    4. Si la fila no tiene valor en nombre_tratamiento, asigna NOMBRE_TRATAMIENTO_DEFECTO.
    5. Celdas combinadas vacías no descartan la fila si contiene datos en otras columnas.

    Retorna (candidatos, total_truncados, total_filas_utiles, total_combinaciones).
    """
    tablas = estructura.get("tablas", [])
    tabla = next((t for t in tablas if t["indice"] == indice_tabla), None)
    if not tabla:
        return [], 0, 0, 0

    filas_texto = tabla.get("filas_texto", [])
    if not filas_texto:
        return [], 0, 0, 0

    encabezados = [c.strip() for c in filas_texto[0]]
    datos_filas = filas_texto[1:] if len(filas_texto) > 1 else []

    candidatos = []
    total_truncados = 0
    total_filas_utiles = 0
    total_combinaciones = 0

    for idx, fila in enumerate(datos_filas):
        fila_limpia = [c.strip() for c in fila]
        if not _es_fila_util(fila_limpia, encabezados):
            continue

        total_filas_utiles += 1
        if len(candidatos) >= limite:
            continue

        nro_fila_original = idx + 2  # 1-based, contando encabezado en fila 1
        campos_candidato = {}
        campos_adicionales = {}

        for idx_col, valor in enumerate(fila_limpia):
            val_limpio, fue_truncado = limitar_texto_campo(valor)
            if fue_truncado:
                total_truncados += 1

            campo_destino = mapeo_columnas.get(idx_col)
            if campo_destino:
                val_existente = campos_candidato.get(campo_destino, "")
                val_combinado = _combinar_valores_campo(val_existente, val_limpio)
                if val_existente and val_limpio and val_existente != val_limpio:
                    total_combinaciones += 1
                campos_candidato[campo_destino] = val_combinado
            elif conservar_adicionales and val_limpio:
                campos_adicionales[f"col_{idx_col}"] = val_limpio

        # Regla del nombre de tratamiento
        if not campos_candidato.get("nombre_tratamiento", "").strip():
            campos_candidato["nombre_tratamiento"] = NOMBRE_TRATAMIENTO_DEFECTO

        datos_json = dict(campos_candidato)
        if conservar_adicionales and campos_adicionales:
            datos_json["campos_adicionales"] = campos_adicionales

        candidatos.append({
            "nro_hoja": f"Tabla {indice_tabla + 1}",
            "nro_fila": nro_fila_original,
            "nro_pagina": None,
            "seccion_origen": f"Tabla {indice_tabla + 1} del documento DOCX",
            "datos_extraidos_json": datos_json,
        })

    return candidatos, total_truncados, total_filas_utiles, total_combinaciones


# ==============================================================================
# PDF CON TEXTO
# ==============================================================================

def extraer_pdf_paginas(contenido: bytes) -> list[dict]:
    """Extrae texto por página del PDF usando pdfplumber."""
    try:
        import pdfplumber
    except ImportError:
        raise ImportError("La biblioteca 'pdfplumber' no está instalada. Ejecute: pip install pdfplumber")

    paginas = []
    try:
        with pdfplumber.open(io.BytesIO(contenido)) as pdf:
            if len(pdf.pages) == 0:
                raise ValueError("El PDF no contiene páginas.")
            for idx, page in enumerate(pdf.pages, start=1):
                try:
                    texto = page.extract_text() or ""
                except Exception:
                    texto = ""
                paginas.append({
                    "nro_pagina": idx,
                    "texto_preview": texto[:600],
                    "texto_completo": texto,
                    "tiene_texto": bool(texto.strip()),
                })
    except Exception as exc:
        raise ValueError(f"No se pudo leer el PDF: {exc}")

    return paginas


def crear_candidato_manual_pdf(
    nro_pagina: int,
    campos_ingresados: dict,
) -> dict:
    """Crea candidato manual desde fragmento de texto PDF."""
    campos_limpios = {}
    for k, v in campos_ingresados.items():
        if v and str(v).strip():
            txt, _ = limitar_texto_campo(str(v).strip())
            campos_limpios[k] = txt

    if not campos_limpios.get("nombre_tratamiento", "").strip():
        campos_limpios["nombre_tratamiento"] = NOMBRE_TRATAMIENTO_DEFECTO

    return {
        "nro_hoja": None,
        "nro_fila": None,
        "nro_pagina": nro_pagina,
        "seccion_origen": f"Página {nro_pagina} del PDF",
        "datos_extraidos_json": campos_limpios,
    }


# ==============================================================================
# MD / TXT
# ==============================================================================

def extraer_texto_plano(contenido: bytes) -> str:
    """Decodifica texto plano (MD o TXT) con fallback seguro."""
    return _decodificar_texto(contenido)


def crear_candidato_manual_texto(
    seccion_origen: str,
    campos_ingresados: dict,
) -> dict:
    """Crea candidato manual desde bloque de texto plano."""
    campos_limpios = {}
    for k, v in campos_ingresados.items():
        if v and str(v).strip():
            txt, _ = limitar_texto_campo(str(v).strip())
            campos_limpios[k] = txt

    if not campos_limpios.get("nombre_tratamiento", "").strip():
        campos_limpios["nombre_tratamiento"] = NOMBRE_TRATAMIENTO_DEFECTO

    return {
        "nro_hoja": None,
        "nro_fila": None,
        "nro_pagina": None,
        "seccion_origen": seccion_origen or "Bloque de texto manual",
        "datos_extraidos_json": campos_limpios,
    }


# ==============================================================================
# RESUMEN EXCLUSIVAMENTE TÉCNICO (PRIVACIDAD Y MINIMIZACIÓN DE DATOS)
# ==============================================================================

def _decodificar_texto(contenido: bytes) -> str:
    """Decodifica bytes a texto con fallback UTF-8 → Latin-1."""
    try:
        return contenido.decode("utf-8")
    except UnicodeDecodeError:
        try:
            return contenido.decode("latin-1")
        except Exception:
            return contenido.decode("utf-8", errors="replace")


def resumen_extraccion(
    formato: str,
    tamano_kb: float,
    total_filas_o_paginas: int,
    columnas_detectadas: list,
    campos_mapeados: list,
    advertencias: list | None = None,
) -> str:
    """Genera resumen estrictamente técnico sin datos personales ni valores de filas."""
    cols_str = ", ".join([str(c)[:30] for c in columnas_detectadas[:6]])
    sufijo_cols = f"… (+{len(columnas_detectadas) - 6})" if len(columnas_detectadas) > 6 else ""

    campos_str = ", ".join(campos_mapeados[:6])
    sufijo_campos = f"… (+{len(campos_mapeados) - 6})" if len(campos_mapeados) > 6 else ""

    partes = [
        f"Tipo: {formato.upper()}",
        f"Tamaño: {tamano_kb:.1f} KB",
        f"Unidades útiles: {total_filas_o_paginas}",
    ]
    if cols_str:
        partes.append(f"Columnas detectadas: [{cols_str}{sufijo_cols}]")
    if campos_str:
        partes.append(f"Campos RAT mapeados: [{campos_str}{sufijo_campos}]")
    if advertencias:
        partes.append(f"Advertencias: {'; '.join(advertencias)}")

    return " | ".join(partes)
